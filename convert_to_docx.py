from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

# Read the markdown file
with open('RESEARCH_REPORT_DETAILED.md', 'r', encoding='utf-8') as f:
    content = f.read()

# Create a Word document
doc = Document()

# Set document margins
sections = doc.sections
for section in sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.25)

# Parse and add content
lines = content.split('\n')

skip_next = 0
i = 0

while i < len(lines):
    if skip_next > 0:
        skip_next -= 1
        i += 1
        continue
    
    line = lines[i]
    
    if not line.strip():
        i += 1
        continue
    
    # Title (# heading)
    if line.startswith('# '):
        title = line.replace('# ', '').strip()
        p = doc.add_paragraph(title, style='Heading 1')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
    # Heading 2 (## heading)
    elif line.startswith('## '):
        heading = line.replace('## ', '').strip()
        doc.add_paragraph(heading, style='Heading 2')
        
    # Heading 3 (### heading)
    elif line.startswith('### '):
        heading = line.replace('### ', '').strip()
        doc.add_paragraph(heading, style='Heading 3')
        
    # Heading 4 (#### heading)
    elif line.startswith('#### '):
        heading = line.replace('#### ', '').strip()
        doc.add_paragraph(heading, style='Heading 4')
        
    # Heading 5 (##### heading)
    elif line.startswith('##### '):
        heading = line.replace('##### ', '').strip()
        doc.add_paragraph(heading, style='Heading 5')
        
    # Table detection
    elif '|' in line and '-' not in line:
        # Check if next line is a separator
        if i + 1 < len(lines) and '|' in lines[i + 1] and '-' in lines[i + 1]:
            # Extract header
            headers = [h.strip() for h in line.split('|')[1:-1]]
            rows = []
            
            # Get all subsequent rows until non-table line
            j = i + 2
            while j < len(lines) and '|' in lines[j] and '-' not in lines[j]:
                cols = [c.strip() for c in lines[j].split('|')[1:-1]]
                if cols and cols[0]:  # If not empty
                    rows.append(cols)
                j += 1
            
            # Create table
            if rows:
                table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
                table.style = 'Light Grid Accent 1'
                
                # Add headers
                hdr_cells = table.rows[0].cells
                for k, header in enumerate(headers):
                    hdr_cells[k].text = header
                    # Bold headers
                    for paragraph in hdr_cells[k].paragraphs:
                        for run in paragraph.runs:
                            run.bold = True
                
                # Add rows
                for row_idx, row_data in enumerate(rows):
                    row_cells = table.rows[row_idx + 1].cells
                    for col_idx, cell_text in enumerate(row_data):
                        if col_idx < len(row_cells):
                            row_cells[col_idx].text = cell_text
            
            skip_next = j - i - 1
    
    # Code blocks
    elif line.startswith('```'):
        # Skip code block lines
        j = i + 1
        while j < len(lines) and not lines[j].startswith('```'):
            j += 1
        skip_next = j - i
        
    # Horizontal rule
    elif line.strip() == '---':
        doc.add_paragraph()
        
    # Regular paragraph
    elif line.strip() and not line.startswith('|'):
        # Process bold and other formatting
        p = doc.add_paragraph()
        
        # Split by ** for bold sections
        parts = re.split(r'\*\*', line)
        for idx, part in enumerate(parts):
            if idx % 2 == 0:  # Regular text
                p.add_run(part)
            else:  # Bold text
                run = p.add_run(part)
                run.bold = True
    
    i += 1

# Save the document
doc.save('RESEARCH_REPORT_DETAILED.docx')
print("✅ Word document created successfully: RESEARCH_REPORT_DETAILED.docx")
